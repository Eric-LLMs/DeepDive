"""Tests for the RAG image pipeline additions: PPTX embedded-image scan + caption chunks.

The PPTX scan keys pictures by 1-based slide number so deck figures ride the same
``[[PAGE:n]]`` anchor axis as PDFs (see ``rag_images`` module docstring); the caption
builder turns each unique image into a searchable leaf chunk with the vision LLM mocked
at the shared ``describe_image`` seam.
"""
from __future__ import annotations

import io

from apps.worker.rag_images import caption_chunks, scan_embedded_images


def _png(color=(255, 0, 0)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), color).save(buf, "PNG")
    return buf.getvalue()


def _pptx_with_pictures(specs: list[list[bytes]]) -> bytes:
    """Build a deck; ``specs[i]`` lists the picture bytes to place on slide i+1."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for pictures in specs:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        for x, blob in enumerate(pictures):
            slide.shapes.add_picture(io.BytesIO(blob), Inches(x), Inches(0), Inches(1), Inches(1))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_scan_pptx_keys_by_slide_number():
    blob = _png()
    data = _pptx_with_pictures([[blob], [], [blob]])
    scans = scan_embedded_images(data, "deck.pptx")
    assert sorted(scans) == [1, 3]  # 1-based slides, empty slide absent
    assert scans[1][0]["data"] == blob
    assert scans[1][0]["mime"] == "image/png"


def test_scan_potx_template_normalizes():
    from tests.test_read_document_tool import _as_alt_package

    blob = _png()
    potx = _as_alt_package(_pptx_with_pictures([[blob]]), "template")
    scans = scan_embedded_images(potx, "lecture.potx")
    assert 1 in scans and scans[1][0]["data"] == blob


def test_scan_pptx_ignores_text_only_decks():
    assert scan_embedded_images(_pptx_with_pictures([[]]), "plain.pptx") == {}


def test_scan_doc_anchors_images_to_pic_paragraphs(monkeypatch):
    """``[pic]`` placeholders from the antiword pass decide the paragraph anchor."""
    from core.infrastructure import doc_images, ingest

    img = {"name": "doc_1.png", "mime": "image/png", "data": b"PNGDATA"}
    monkeypatch.setattr(doc_images, "scan_doc_images", lambda data: ([img], 0))
    monkeypatch.setattr(
        ingest, "extract_text",
        lambda data, name, *, para_markers=False:
            "[[PARA:0]]\nintro\n\n[[PARA:3]]\nsee [pic] here",
    )
    scans = scan_embedded_images(b"container", "old.doc")
    assert scans == {3: [img]}


def test_scan_doc_falls_back_to_ordinal_anchors_without_text_pass(monkeypatch):
    """No antiword on the box → anchoring degrades, but images are still scanned."""
    from core.infrastructure import doc_images, ingest
    from core.infrastructure.ingest import UnsupportedFileType

    imgs = [
        {"name": "doc_1.png", "mime": "image/png", "data": b"A"},
        {"name": "doc_2.jpg", "mime": "image/jpeg", "data": b"B"},
    ]
    monkeypatch.setattr(doc_images, "scan_doc_images", lambda data: (imgs, 1))

    def _boom(data, name, *, para_markers=False):
        raise UnsupportedFileType("antiword not installed")

    monkeypatch.setattr(ingest, "extract_text", _boom)
    scans = scan_embedded_images(b"container", "old.doc")
    assert scans == {0: [imgs[0]], 1: [imgs[1]]}


def test_scan_doc_no_images_returns_empty(monkeypatch):
    from core.infrastructure import doc_images

    monkeypatch.setattr(doc_images, "scan_doc_images", lambda data: ([], 0))
    assert scan_embedded_images(b"container", "old.doc") == {}


def test_extract_ppt_page_markers_roundtrip():
    from apps.worker.tasks import _MARKER_STRIP, _PAGE_MARKER
    from core.infrastructure.ingest import extract_text

    data = _pptx_with_pictures([[_png()], []])
    text = extract_text(data, "deck.pptx", para_markers=True)
    assert _PAGE_MARKER.findall(text) == ["1", "2"]  # every slide anchored, 1-based
    assert "[[PAGE:" not in _MARKER_STRIP.sub("", text)


async def test_caption_chunks_dedupes_by_bytes(monkeypatch):
    import core.infrastructure.vision_caption as vc

    calls: list[bytes] = []
    uids: list = []

    async def _d(data, mime="", *, llm, session_factory, prompt=None, user_id=None, role_id=None):
        calls.append(data)
        uids.append(user_id)
        return "a red square"

    monkeypatch.setattr(vc, "describe_image", _d)
    blob = _png()
    scans = {1: [{"name": "a.png", "mime": "image/png", "data": blob}],
             2: [{"name": "b.png", "mime": "image/png", "data": blob}]}
    ids = {1: ["id-1"], 2: ["id-2"]}
    chunks = await caption_chunks(
        scans, ids, "deck.pptx",
        llm=None, session_factory=None, axis_key="pages", anchor_label="slide",
        user_id="owner-1",
    )
    assert len(calls) == 1  # same bytes captioned once
    assert uids == ["owner-1"]  # owner identity reaches the funnel
    assert len(chunks) == 1
    c = chunks[0]
    assert "slide 1, slide 2" in c.content_en and "a red square" in c.content_en
    assert c.meta["kind"] == "image_caption"
    assert c.meta["image_id"] == "id-1"
    assert c.meta["image_ids"] == ["id-1", "id-2"]
    assert c.meta["pages"] == [1, 2]


async def test_caption_chunks_skip_failures(monkeypatch):
    import core.infrastructure.vision_caption as vc

    async def _boom(data, mime="", *, llm, session_factory, prompt=None, user_id=None, role_id=None):
        raise RuntimeError("vision down")

    monkeypatch.setattr(vc, "describe_image", _boom)
    scans = {1: [{"name": "a.png", "mime": "image/png", "data": _png()}]}
    chunks = await caption_chunks(
        scans, {1: ["id-1"]}, "doc.pdf",
        llm=None, session_factory=None, axis_key="pages", anchor_label="page",
    )
    assert chunks == []  # ingest survives caption failure


async def test_caption_chunks_empty_caption_skipped(monkeypatch):
    import core.infrastructure.vision_caption as vc

    async def _empty(data, mime="", *, llm, session_factory, prompt=None, user_id=None, role_id=None):
        return "   "

    monkeypatch.setattr(vc, "describe_image", _empty)
    scans = {3: [{"name": "a.png", "mime": "image/png", "data": _png()}]}
    chunks = await caption_chunks(
        scans, {3: ["x"]}, "doc.pdf",
        llm=None, session_factory=None, axis_key="pages", anchor_label="page",
    )
    assert chunks == []


def test_pipeline_config_image_captions_default_on():
    from rag.pipeline.pipeline_config import RagPipelineConfig

    cfg = RagPipelineConfig.from_dict({})
    assert cfg.image_captions is True  # stored configs predating the flag opt in
    assert RagPipelineConfig.from_dict(cfg.to_dict()).image_captions is True
    assert RagPipelineConfig.from_dict({"image_captions": False}).image_captions is False
    assert RagPipelineConfig.from_dict(
        RagPipelineConfig(image_captions=False).to_dict()
    ).image_captions is False


def test_vision_tool_reuses_shared_channel_helper():
    """The refactor moved channel resolution to core; the tool module must still register."""
    import apps.api.tools.vision_tool as vt
    from agent import ToolRuntime

    assert vt.describe_image.__module__ == "core.infrastructure.vision_caption"
    vt.register(ToolRuntime(), None, llm=None)  # registration side effect must not raise
