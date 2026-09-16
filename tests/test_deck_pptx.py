"""Native PptxCompiler (M2): package integrity, per-template shapes, loud fails.

The all-templates golden brief doubles as the PPTX fixture: every branch of
:mod:`deck.compiler.pptx_builder` must draw, the package must reopen with
python-pptx, text must stay in native editable objects, charts must embed the
SAME materializer bytes as the shared PNG source, and an unmapped template
must fail loudly instead of flattening to bullets (§9.4).
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from apps.api.tools.toolkit.deck import schema as S
from apps.api.tools.toolkit.deck.compiler import layout_engine as LE
from apps.api.tools.toolkit.deck.compiler import pptx_builder as PB
from apps.api.tools.toolkit.deck.compiler.materializer import render_chart_png
from apps.api.tools.toolkit.deck.compiler.theme import DARK_BLUEPRINT, theme_for
from apps.api.tools.toolkit.deck.errors import DeckLayoutError
from apps.api.tools.toolkit.deck.layout import PAGE_H_MM, PAGE_W_MM
from apps.api.tools.toolkit.deck.render import brief_to_pptx
from tests.test_deck_compiler import brief_from, slide
from tests.test_deck_typst_golden import all_templates_brief


@pytest.fixture()
def deck(tmp_path):
    brief, _layouts, asset = all_templates_brief(tmp_path)
    return brief, [asset]


def reopened(data: bytes) -> Presentation:
    assert data[:2] == b"PK"                    # a real OOXML package
    return Presentation(BytesIO(data))


def texts(slide) -> list[str]:
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                out += [c.text for c in row.cells]
    return out


class TestPackage:
    def test_page_count_and_aspect(self, deck):
        prs = reopened(brief_to_pptx(*deck, document_title="RAG Survey"))
        slides = list(prs.slides)
        assert len(slides) == 1 + len(deck[0].slides)
        assert abs(prs.slide_width / prs.slide_height
                   - PAGE_W_MM / PAGE_H_MM) < 0.01

    def test_dispatch_covers_every_happy_fn(self):
        fns = {fn for _, fn in LE.HAPPY.values()}
        assert fns <= set(PB._DISPATCH)

    def test_unmapped_fn_fails_loudly(self, deck):
        brief, _assets = deck
        bogus = [LE.BriefSlideLayout(slide_index=1, fn="spaceshipSlide",
                                     template=LE.LayoutTemplate.CENTERED_THESIS)]
        one = brief.model_copy(update={"slides": brief.slides[:1]})
        with pytest.raises(DeckLayoutError, match="no PPTX branch"):
            PB.build_brief_pptx(one, bogus, {}, theme=theme_for("x"))

    def test_layout_count_mismatch_fails(self, deck):
        brief, _assets = deck
        with pytest.raises(DeckLayoutError, match="layouts"):
            PB.build_brief_pptx(brief, [], {}, theme=theme_for("x"))

    def test_structural_stability_across_builds(self, deck):
        a = reopened(brief_to_pptx(*deck, document_title="T"))
        b = reopened(brief_to_pptx(*deck, document_title="T"))
        la, lb = list(a.slides), list(b.slides)
        assert len(la) == len(lb)
        for sa, sb in zip(la, lb):
            assert len(sa.shapes) == len(sb.shapes)
            assert texts(sa) == texts(sb)


class TestContent:
    def _slides(self, deck):
        prs = reopened(brief_to_pptx(*deck, document_title="RAG Survey",
                                     source_names=["doc1.pdf"]))
        return brief_to_map(prs, deck[0])

    def test_titles_are_native_text(self, deck):
        by_idx = self._slides(deck)
        for plan in deck[0].slides:
            assert any(plan.title in t for t in texts(by_idx[plan.slide_index]))

    def test_card_text_survives_verbatim(self, deck):
        by_idx = self._slides(deck)
        for plan in deck[0].slides:
            for c in plan.cards:
                blob = " ".join(texts(by_idx[plan.slide_index]))
                assert c.label in blob and c.takeaway in blob

    def test_notes_and_citations(self, deck):
        by_idx = self._slides(deck)
        t1 = by_idx[1]
        assert t1.has_notes_slide
        assert "notes" in t1.notes_slide.notes_text_frame.text.lower()
        carded = by_idx[2]                           # STRUCTURED_CARDS cites t1
        assert "[doc1:3]" in carded.notes_slide.notes_text_frame.text
        blob = " ".join(texts(carded))
        assert "[doc1:3]" in blob                   # and on the slide face, like the PDF

    def test_table_grammars_are_real_tables(self, deck):
        by_idx = self._slides(deck)
        for plan in deck[0].slides:
            if plan.visual_spec.grammar.value in ("TABLE", "COMPARISON"):
                kinds = {sh.shape_type for sh in by_idx[plan.slide_index].shapes}
                assert MSO_SHAPE_TYPE.TABLE in kinds

    def test_chart_png_is_the_shared_materializer_bytes(self, deck):
        brief, _assets = deck
        by_idx = self._slides(deck)
        chart = next(p for p in brief.slides
                     if p.visual_spec.grammar is S.VisualGrammar.DATA_CHART)
        kind, labels, values, _ = LE.chart_frame(chart.visual_spec.generation_spec)
        want = render_chart_png(labels, values, kind=kind, name=chart.title,
                                theme=theme_for(brief.presentation_style))
        pics = [sh for sh in by_idx[chart.slide_index].shapes
                if sh.shape_type is MSO_SHAPE_TYPE.PICTURE]
        assert len(pics) == 1
        assert pics[0].image.blob == want

    def test_figure_embeds_asset_verbatim(self, deck):
        brief, assets = deck
        by_idx = self._slides(deck)
        fig = next(p for p in brief.slides
                   if p.visual_spec.grammar is S.VisualGrammar.SOURCE_FIGURE_REUSE)
        pics = [sh for sh in by_idx[fig.slide_index].shapes
                if sh.shape_type is MSO_SHAPE_TYPE.PICTURE]
        assert len(pics) == 1
        assert pics[0].image.blob == Path(assets[0].path).read_bytes()

    def test_dark_theme_reaches_the_cover(self, tmp_path):
        brief, _layouts, asset = all_templates_brief(tmp_path)
        dark = brief.model_copy(update={"presentation_style": "Dark Blueprint"})
        prs = reopened(brief_to_pptx(dark, [asset]))
        theme = theme_for("Dark Blueprint")
        assert theme is DARK_BLUEPRINT
        bg = next(iter(prs.slides)).shapes[0]
        assert str(bg.fill.fore_color.rgb) == "{:02X}{:02X}{:02X}".format(
            *theme.background)


def brief_to_map(prs: Presentation, brief) -> dict[int, object]:
    """slide_index → slide object (index 0 is the cover)."""
    slides = list(prs.slides)
    return {plan.slide_index: slides[i + 1]
            for i, plan in enumerate(brief.slides)}


# ── degraded briefs still compile (thesis/cards fallback legs) ────────────────

def test_degraded_slides_build(tmp_path):
    brief = brief_from([slide(1, "SYSTEM_BLUEPRINT", n_cards=0)])
    data = brief_to_pptx(brief, [])
    prs = reopened(data)
    assert len(list(prs.slides)) == 2
